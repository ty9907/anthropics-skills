"""
仓库哨兵（Repo Sentinel）- 节点三辅助：写回分析结果并生成最终分析报告

将 sub-agent 的分析结果写回单仓库扫描报告（XLSX），
同时生成该仓库的最终分析报告，包含：
- 敏感信息确认结果
- 风险等级评估
- 处理建议
- 仓库信息标识

分析仅基于扫描报告中的匹配代码片段，不依赖原始代码文件。

使用方式：
    # 单条写回
    python -m scripts.write_analysis --report ./workspace/reports/repo1_scan.xlsx \
        --index 0 --result '{"is_sensitive":"是","confirmed_risk_level":"High","llm_reasoning":"...","recommendation":"..."}'

    # 批量写回
    python -m scripts.write_analysis --report ./workspace/reports/repo1_scan.xlsx \
        --results-file ./workspace/analysis_results.json
"""

import argparse
import json
import sys
from collections import defaultdict
from pathlib import Path

from utils import iso_now, write_status


ANALYSIS_COLUMNS = ["is_sensitive", "confirmed_risk_level", "llm_reasoning", "recommendation", "analysis_timestamp"]


def write_single_result(report_path: str, index: int, result: dict) -> None:
    """
    将单条分析结果写回扫描报告

    参数：
        report_path: 扫描报告路径
        index: Finding 索引（从0开始，对应 Findings 工作表第2行起）
        result: 分析结果字典
    """
    from openpyxl import load_workbook

    wb = load_workbook(report_path)
    ws = wb["Findings"]

    headers = [cell.value for cell in ws[1]]

    for col_name in ANALYSIS_COLUMNS:
        if col_name not in headers:
            headers.append(col_name)
            col_idx = len(headers)
            ws.cell(row=1, column=col_idx, value=col_name)

    col_map = {name: idx + 1 for idx, name in enumerate(headers)}

    target_row = index + 2
    if target_row > ws.max_row:
        print(f"警告：索引 {index} 超出范围（最大行号 {ws.max_row - 1}）", file=sys.stderr)
        wb.close()
        return

    for col_name in ANALYSIS_COLUMNS:
        if col_name in col_map:
            value = result.get(col_name, "")
            if col_name == "analysis_timestamp" and not value:
                value = iso_now()
            ws.cell(row=target_row, column=col_map[col_name], value=value)

    wb.save(report_path)
    wb.close()
    print(f"已写入第 {index + 1} 条发现的分析结果")


def write_batch_results(report_path: str, results: list[dict]) -> None:
    """
    将批量分析结果写回扫描报告

    参数：
        report_path: 扫描报告路径
        results: 分析结果列表，每项包含 index 和分析字段
    """
    from openpyxl import load_workbook

    wb = load_workbook(report_path)
    ws = wb["Findings"]

    headers = [cell.value for cell in ws[1]]

    for col_name in ANALYSIS_COLUMNS:
        if col_name not in headers:
            headers.append(col_name)
            col_idx = len(headers)
            ws.cell(row=1, column=col_idx, value=col_name)

    col_map = {name: idx + 1 for idx, name in enumerate(headers)}

    written = 0
    for result in results:
        index = result.get("index")
        if index is None:
            print(f"警告：跳过缺少 index 字段的结果: {result}", file=sys.stderr)
            continue

        target_row = int(index) + 2
        if target_row > ws.max_row:
            print(f"警告：索引 {index} 超出范围，跳过", file=sys.stderr)
            continue

        for col_name in ANALYSIS_COLUMNS:
            if col_name in col_map:
                value = result.get(col_name, "")
                if col_name == "analysis_timestamp" and not value:
                    value = iso_now()
                ws.cell(row=target_row, column=col_map[col_name], value=value)

        written += 1

    wb.save(report_path)
    wb.close()
    print(f"已写入 {written}/{len(results)} 条分析结果")


def generate_final_report(scan_report_path: str, output_path: str) -> None:
    """
    基于已写入分析结果的扫描报告，生成最终分析报告

    最终分析报告包含：
    - Findings：所有匹配记录及分析确认结果
    - Confirmed Sensitive：确认为敏感信息的记录
    - By Risk Level：按确认风险等级分组
    - By Category：按规则分类分组
    - Summary：统计汇总（含仓库名称、确认统计、处理建议汇总）

    参数：
        scan_report_path: 已写入分析结果的扫描报告路径
        output_path: 最终分析报告输出路径
    """
    from openpyxl import load_workbook, Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb_src = load_workbook(scan_report_path, read_only=True)
    ws_src = wb_src["Findings"]

    headers = [cell.value for cell in ws_src[1]]
    all_rows = []
    for row in ws_src.iter_rows(min_row=2):
        row_dict = {}
        for i, cell in enumerate(row):
            if i < len(headers) and headers[i]:
                row_dict[headers[i]] = cell.value
        if row_dict:
            all_rows.append(row_dict)

    wb_src.close()

    repo_name = ""
    if all_rows:
        repo_name = str(all_rows[0].get("repo_name", ""))

    wb = Workbook()

    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font_white = Font(bold=True, color="FFFFFF")
    risk_fill = {
        "High": PatternFill(start_color="FF4444", end_color="FF4444", fill_type="solid"),
        "Medium": PatternFill(start_color="FFA500", end_color="FFA500", fill_type="solid"),
        "Low": PatternFill(start_color="44AA44", end_color="44AA44", fill_type="solid"),
    }

    all_columns = headers + [c for c in ANALYSIS_COLUMNS if c not in headers]

    ws_findings = wb.active
    ws_findings.title = "Findings"
    ws_findings.append(all_columns)
    for cell in ws_findings[1]:
        cell.font = header_font_white
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for row_dict in all_rows:
        ws_findings.append([row_dict.get(col, "") for col in all_columns])
    for col_idx in range(1, len(all_columns) + 1):
        col_letter = ws_findings.cell(row=1, column=col_idx).column_letter
        ws_findings.column_dimensions[col_letter].width = 20

    if "confirmed_risk_level" in all_columns:
        risk_col_idx = all_columns.index("confirmed_risk_level") + 1
        for row in ws_findings.iter_rows(min_row=2, min_col=risk_col_idx, max_col=risk_col_idx):
            for cell in row:
                if cell.value in risk_fill:
                    cell.fill = risk_fill[cell.value]
                    cell.font = Font(color="FFFFFF", bold=True)

    confirmed_rows = [r for r in all_rows if r.get("is_sensitive") == "是"]
    if confirmed_rows:
        ws_confirmed = wb.create_sheet("Confirmed Sensitive")
        ws_confirmed.append(all_columns)
        for cell in ws_confirmed[1]:
            cell.font = header_font_white
            cell.fill = PatternFill(start_color="FF4444", end_color="FF4444", fill_type="solid")
            cell.alignment = Alignment(horizontal="center")
        for row_dict in confirmed_rows:
            ws_confirmed.append([row_dict.get(col, "") for col in all_columns])
        for col_idx in range(1, len(all_columns) + 1):
            col_letter = ws_confirmed.cell(row=1, column=col_idx).column_letter
            ws_confirmed.column_dimensions[col_letter].width = 20

    by_risk = defaultdict(list)
    by_category = defaultdict(list)
    for r in all_rows:
        confirmed_level = r.get("confirmed_risk_level", "Unknown")
        if r.get("is_sensitive") != "是":
            confirmed_level = "N/A"
        by_risk[confirmed_level].append(r)
        by_category[r.get("category", "unknown")].append(r)

    def _write_pivot(ws, groups, key_name):
        ws.append([key_name, "Count", "High", "Medium", "Low", "N/A"])
        for cell in ws[1]:
            cell.font = header_font_white
            cell.fill = header_fill
        for key, items in sorted(groups.items()):
            high = sum(1 for i in items if i.get("confirmed_risk_level") == "High")
            medium = sum(1 for i in items if i.get("confirmed_risk_level") == "Medium")
            low = sum(1 for i in items if i.get("confirmed_risk_level") == "Low")
            na = sum(1 for i in items if i.get("is_sensitive") != "是")
            ws.append([key, len(items), high, medium, low, na])

    ws_risk = wb.create_sheet("By Risk Level")
    _write_pivot(ws_risk, by_risk, "Confirmed Risk Level")

    ws_category = wb.create_sheet("By Category")
    _write_pivot(ws_category, by_category, "Category")

    ws_summary = wb.create_sheet("Summary")
    ws_summary.append(["Metric", "Value"])
    ws_summary.append(["Repository", repo_name])
    ws_summary.append(["Total Findings", len(all_rows)])
    ws_summary.append(["Confirmed Sensitive", len(confirmed_rows)])
    ws_summary.append(["Rejected (Not Sensitive)", sum(1 for r in all_rows if r.get("is_sensitive") == "否")])
    ws_summary.append(["Unknown", sum(1 for r in all_rows if r.get("is_sensitive") not in ("是", "否"))])
    ws_summary.append(["High Risk", len(by_risk.get("High", []))])
    ws_summary.append(["Medium Risk", len(by_risk.get("Medium", []))])
    ws_summary.append(["Low Risk", len(by_risk.get("Low", []))])
    ws_summary.append(["Unique Files Affected", len(set(r.get("file_path") for r in confirmed_rows))])
    ws_summary.append(["Analysis Completed", iso_now()])

    if confirmed_rows:
        ws_summary.append([])
        ws_summary.append(["处理建议汇总", ""])
        recommendations = set()
        for r in confirmed_rows:
            rec = r.get("recommendation", "")
            if rec and str(rec).strip():
                recommendations.add(str(rec).strip())
        for i, rec in enumerate(sorted(recommendations), 1):
            ws_summary.append([f"建议 {i}", rec])

    for cell in ws_summary[1]:
        cell.font = header_font_white
        cell.fill = header_fill

    wb.save(output_path)
    print(f"最终分析报告已生成: {output_path}")


def main(args: list[str] | None = None) -> None:
    """
    主入口：将分析结果写回扫描报告并生成最终分析报告
    """
    parser = argparse.ArgumentParser(description="写回分析结果并生成最终分析报告")
    parser.add_argument("--report", required=True, help="单仓库扫描报告 XLSX 文件路径")
    parser.add_argument("--index", type=int, help="单条模式：Finding 索引（从0开始）")
    parser.add_argument("--result", help="单条模式：分析结果 JSON 字符串")
    parser.add_argument("--results-file", help="批量模式：分析结果 JSON 文件路径")
    parser.add_argument("--workspace", help="工作目录（用于写状态文件和生成最终报告）")

    parsed = parser.parse_args(args)
    report_path = Path(parsed.report)

    if not report_path.exists():
        print(f"错误：报告文件不存在: {report_path}", file=sys.stderr)
        sys.exit(1)

    if parsed.index is not None and parsed.result:
        try:
            result = json.loads(parsed.result)
        except json.JSONDecodeError as e:
            print(f"错误：无法解析结果 JSON: {e}", file=sys.stderr)
            sys.exit(1)

        write_single_result(str(report_path), parsed.index, result)

        if parsed.workspace:
            status_path = Path(parsed.workspace) / "analysis_status.json"
            write_status(status_path, "in_progress", items_total=1, items_completed=1, items_failed=0)

    elif parsed.results_file:
        results_path = Path(parsed.results_file)
        if not results_path.exists():
            print(f"错误：结果文件不存在: {results_path}", file=sys.stderr)
            sys.exit(1)

        try:
            results = json.loads(results_path.read_text(encoding="utf-8"))
        except (json.JSONDecodeError, OSError) as e:
            print(f"错误：无法读取结果文件: {e}", file=sys.stderr)
            sys.exit(1)

        if not isinstance(results, list):
            results = [results]

        write_batch_results(str(report_path), results)

        if parsed.workspace:
            status_path = Path(parsed.workspace) / "analysis_status.json"
            confirmed = sum(1 for r in results if r.get("is_sensitive") == "是")
            rejected = sum(1 for r in results if r.get("is_sensitive") == "否")
            unknown = sum(1 for r in results if r.get("is_sensitive") not in ("是", "否"))
            final_status = "completed" if unknown == 0 else "partial"
            write_status(
                status_path, final_status,
                items_total=len(results),
                items_completed=confirmed + rejected,
                items_failed=unknown,
            )
    else:
        print("错误：请指定 --index + --result（单条模式）或 --results-file（批量模式）", file=sys.stderr)
        parser.print_help()
        sys.exit(1)

    if parsed.workspace:
        reports_dir = Path(parsed.workspace) / "reports"
        reports_dir.mkdir(parents=True, exist_ok=True)

        repo_name = report_path.stem.replace("_scan", "")
        final_report_path = str(reports_dir / f"{repo_name}_analysis.xlsx")
        generate_final_report(str(report_path), final_report_path)

        if parsed.workspace:
            status_path = Path(parsed.workspace) / "analysis_status.json"
            write_status(status_path, "completed", final_report=final_report_path)

    print(f"扫描报告已更新: {report_path}")
